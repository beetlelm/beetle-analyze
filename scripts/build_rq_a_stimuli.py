#!/usr/bin/env python3
"""
scripts/build_rq_a_stimuli.py — Build RQ-A German↔English cross-lingual priming
stimuli from real, published items in van Dijk & Hopp (2025).

Source paper
------------
van Dijk, C., & Hopp, H. (2025). Structural preferences in language learning
across languages. *Language, Cognition and Neuroscience*.
DOI: 10.1080/23273798.2025.2476703

Source materials (open data)
----------------------------
OSF project ypsvn (https://osf.io/ypsvn). Materials directory contains
``WP4_stimuli_OSF_20240826.xlsx`` (~70 KB), with per-item bilingual sentence
pairs for two construction types:

  * Sheet ``Ditransitive_prime_target_pairs`` (32 prime/target pairs):
    German prime sentences in PO and DO form, English target sentences in
    PO and DO form. Includes a ``Category_prime`` column (``strong_DO`` /
    ``weak_DO``) capturing the verb-bias factor that is the paper's headline
    contribution.

  * Sheet ``Transitives_prime_target_pairs`` (16 prime/target pairs):
    German prime sentences in active and passive form, English target
    sentences in active and passive form. German passives use werden + von /
    vom + past participle.

In both sheets every pair has Language_prime='german', Language_target=
'english' — i.e. the experimental design tests deu→eng priming only. We
preserve that direction faithfully; we do not synthesise an eng→deu
direction from these items because the picture-matching design is
direction-specific.

Output
------
Two schema-B TSVs at ``arnett-2025-priming-raw/stimuli/``:

  * ``vandijk_hopp_2025_german_dat.tsv`` — dative DO/PO priming
    (32 pairs × 4 type combos = 128 data rows)

  * ``vandijk_hopp_2025_german_actpas.tsv`` — active/passive priming
    (16 pairs × 4 type combos = 64 data rows)

Schema (matches Arnett 2025 ``more_*_stims_144.tsv`` files)::

  <unnamed-index>\\tSet\\tPrimeLanguage\\tTargetLanguage\\tPrimeType\\t
  TargetType\\tPrime\\tTarget\\tCategory

The trailing ``Category`` column carries the strong_DO / weak_DO tag for
ditransitives, and is empty for transitives. ``evals/priming_stimuli.py``
ignores any extra columns, so this is forward-compatible with the existing
loader.

Usage
-----
::

    python scripts/build_rq_a_stimuli.py

The script will auto-download the xlsx from OSF if it is not present locally
under ``arnett-2025-priming-raw/stimuli/source/``. No HuggingFace or OSF
auth is required — the file is openly accessible.

Citation policy
---------------
The stimulus sentences belong to van Dijk & Hopp (2025), distributed under
the OSF project's CC license. This script does not modify the sentences;
it only re-shapes them into the schema the Beetle priming pipeline expects.
Any publication using the resulting TSVs should cite the source paper.
"""
from __future__ import annotations

import csv
import sys
import urllib.request
from itertools import product
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
STIMULI_DIR = REPO_ROOT / "arnett-2025-priming-raw" / "stimuli"
SOURCE_DIR = STIMULI_DIR / "source"
XLSX_NAME = "WP4_stimuli_OSF_20240826.xlsx"
XLSX_PATH = SOURCE_DIR / XLSX_NAME

# OSF download URL for the materials xlsx (file guid 3bgjn under project ypsvn).
# This is the openly-shared file from the OSF project; no auth required.
# Mirror via the waterbutler endpoint is kept as a fallback in case the GUID
# redirect ever changes — both resolve to the same stored file.
OSF_XLSX_URL = "https://osf.io/download/3bgjn/"
OSF_XLSX_URL_FALLBACK = (
    "https://files.de-1.osf.io/v1/resources/ypsvn/providers/osfstorage/"
    "66cef586cad2f334508c90b3"
)

OUT_DAT = STIMULI_DIR / "vandijk_hopp_2025_german_dat.tsv"
OUT_ACTPAS = STIMULI_DIR / "vandijk_hopp_2025_german_actpas.tsv"


def _download(url: str, dest: Path) -> None:
    dest.parent.mkdir(parents=True, exist_ok=True)
    print(f"[osf] downloading {url} -> {dest}")
    req = urllib.request.Request(url, headers={"User-Agent": "beetle-analyze/1.0"})
    with urllib.request.urlopen(req, timeout=60) as resp, dest.open("wb") as f:
        f.write(resp.read())
    print(f"[osf] wrote {dest.stat().st_size:,} bytes")


def _ensure_xlsx() -> Path:
    if XLSX_PATH.exists() and XLSX_PATH.stat().st_size > 10_000:
        # Quick zip-magic sanity check; OSF occasionally serves an HTML error
        # page through the GUID redirect, which would land here as a non-zip.
        with XLSX_PATH.open("rb") as f:
            magic = f.read(2)
        if magic == b"PK":
            return XLSX_PATH
        print(f"[osf] cached xlsx is not a valid zip — re-downloading")
        XLSX_PATH.unlink()
    for url in (OSF_XLSX_URL, OSF_XLSX_URL_FALLBACK):
        try:
            _download(url, XLSX_PATH)
            with XLSX_PATH.open("rb") as f:
                if f.read(2) == b"PK":
                    return XLSX_PATH
            print(f"[osf] {url} returned non-zip content — trying next mirror")
        except Exception as e:
            print(f"[osf] {url} failed: {e}")
    raise RuntimeError("Could not fetch a valid xlsx from OSF mirrors")


def _load_sheet(path: Path, sheet: str) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        print("openpyxl is required: pip install openpyxl", file=sys.stderr)
        raise
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(c) if c is not None else "" for c in rows[0]]
    out = []
    for r in rows[1:]:
        d = {h: v for h, v in zip(headers, r)}
        if d.get("ID_prime"):
            out.append(d)
    return out


def _lang_title(s: str) -> str:
    """Map xlsx language codes to bare title-case names matching schema B."""
    return {"english": "English", "german": "German"}[str(s).strip().lower()]


def _build_dative_rows(pairs: list[dict]) -> list[list]:
    """For each (prime, target) pair, produce 4 rows = PO/DO × PO/DO."""
    out = []
    idx = 0
    for p in pairs:
        set_id = f"{p['ID_prime']}__{p['ID_target']}"
        prime_lang = _lang_title(p["Language_prime"])
        target_lang = _lang_title(p["Language_target"])
        prime_sents = {"PO": p["PO_prime"], "DO": p["DO_prime"]}
        target_sents = {"PO": p["PO_target"], "DO": p["DO_target"]}
        category = p.get("Category_prime", "") or ""
        for pt, tt in product(("PO", "DO"), ("PO", "DO")):
            out.append([
                idx, set_id, prime_lang, target_lang, pt, tt,
                prime_sents[pt], target_sents[tt], category,
            ])
            idx += 1
    return out


def _build_actpas_rows(pairs: list[dict]) -> list[list]:
    """For each (prime, target) pair, produce 4 rows = active/passive × active/passive."""
    out = []
    idx = 0
    for p in pairs:
        set_id = f"{p['ID_prime']}__{p['ID_target']}"
        prime_lang = _lang_title(p["Language_prime"])
        target_lang = _lang_title(p["Language_target"])
        prime_sents = {"active": p["Active_prime"], "passive": p["Passive_prime"]}
        target_sents = {"active": p["Active_target"], "passive": p["Passive_target"]}
        for pt, tt in product(("active", "passive"), ("active", "passive")):
            out.append([
                idx, set_id, prime_lang, target_lang, pt, tt,
                prime_sents[pt], target_sents[tt], "",
            ])
            idx += 1
    return out


def _write_tsv(path: Path, rows: list[list]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    header = ["", "Set", "PrimeLanguage", "TargetLanguage",
              "PrimeType", "TargetType", "Prime", "Target", "Category"]
    with path.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, delimiter="\t", quoting=csv.QUOTE_MINIMAL)
        w.writerow(header)
        for r in rows:
            w.writerow(r)
    print(f"[write] {path}  ({len(rows)} rows)")


def main() -> None:
    xlsx = _ensure_xlsx()
    dat_pairs = _load_sheet(xlsx, "Ditransitive_prime_target_pairs")
    actpas_pairs = _load_sheet(xlsx, "Transitives_prime_target_pairs")
    print(f"[parse] {len(dat_pairs)} dative pairs, {len(actpas_pairs)} actpas pairs")

    dat_rows = _build_dative_rows(dat_pairs)
    actpas_rows = _build_actpas_rows(actpas_pairs)
    assert len(dat_rows) == len(dat_pairs) * 4, len(dat_rows)
    assert len(actpas_rows) == len(actpas_pairs) * 4, len(actpas_rows)

    _write_tsv(OUT_DAT, dat_rows)
    _write_tsv(OUT_ACTPAS, actpas_rows)


if __name__ == "__main__":
    main()
